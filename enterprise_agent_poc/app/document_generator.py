"""Activity Plan DOCX generation with presentation-only rendering.

The content contract is owned by the Activity Plan Skill.  This module never
calls a model and never rewrites content; it validates bounded input, renders a
DOCX, and stores the immutable result through the existing object-store layer.
"""
from __future__ import annotations

import io
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import PurePosixPath
from typing import Annotated, Callable
from uuid import uuid4

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

from app.brand_logo import BrandLogoError, inspect_transparent_png, validate_brand_logo_key
from app.storage import StorageObjectNotFound, StorageProvider, StorageUnavailable


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_DOCUMENT_BYTES = 12 * 1024 * 1024
MAX_LOGO_BYTES = 2 * 1024 * 1024
TEMPLATE_LOGGER = logging.getLogger(__name__)
DOCUMENT_ID_RE = re.compile(r"^[0-9a-f]{32}$")
WINDOWS_ILLEGAL_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
INVALID_XML_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))
}

ShortText = Annotated[str, Field(max_length=500)]
ListText = Annotated[str, Field(max_length=1_000)]


class ActivityPlanItem(BaseModel):
    model_config = ConfigDict(extra="forbid")

    phase: Annotated[str, Field(max_length=120)]
    name: Annotated[str, Field(max_length=200)]
    description: Annotated[str, Field(max_length=4_000)]
    image_requirement: Annotated[str, Field(max_length=1_000)]

    @field_validator("phase", "name", "description", "image_requirement")
    @classmethod
    def valid_xml_text(cls, value: str) -> str:
        if INVALID_XML_CONTROL.search(value):
            raise ValueError("文档内容包含不支持的控制字符。")
        return value


class ActivityPlanContent(BaseModel):
    """The existing Activity Plan Structured Result Contract, unchanged."""

    model_config = ConfigDict(extra="forbid")

    document_title: Annotated[str, Field(max_length=200)]
    activity_theme: ShortText
    activity_time: ShortText
    activity_location: ShortText
    target_audience: Annotated[str, Field(max_length=1_000)]
    promotion_channels: list[ListText] = Field(max_length=30)
    activity_items: list[ActivityPlanItem] = Field(max_length=80)
    invitation_copy: Annotated[str, Field(max_length=20_000)]
    pending_items: list[ListText] = Field(max_length=80)

    @field_validator(
        "document_title", "activity_theme", "activity_time", "activity_location", "target_audience", "invitation_copy"
    )
    @classmethod
    def valid_xml_text(cls, value: str) -> str:
        if INVALID_XML_CONTROL.search(value):
            raise ValueError("文档内容包含不支持的控制字符。")
        return value

    @field_validator("promotion_channels", "pending_items")
    @classmethod
    def valid_xml_list(cls, values: list[str]) -> list[str]:
        if any(INVALID_XML_CONTROL.search(value) for value in values):
            raise ValueError("文档内容包含不支持的控制字符。")
        return values

    @model_validator(mode="after")
    def bounded_total_content(self) -> "ActivityPlanContent":
        values = [
            self.document_title, self.activity_theme, self.activity_time, self.activity_location,
            self.target_audience, self.invitation_copy, *self.promotion_channels, *self.pending_items,
        ]
        values.extend(
            value
            for item in self.activity_items
            for value in (item.phase, item.name, item.description, item.image_requirement)
        )
        if sum(len(value) for value in values) > 80_000:
            raise ValueError("活动方案内容总长度超过 DOCX V1 限制。")
        return self


class DocumentPresentationOptions(BaseModel):
    model_config = ConfigDict(extra="forbid")

    brand_name: Annotated[str, Field(max_length=120)] = ""
    logo_path: Annotated[str, Field(max_length=512)] = ""
    image_paths: dict[int, Annotated[str, Field(max_length=512)]] = Field(default_factory=dict, max_length=80)
    primary_color: Annotated[str, Field(pattern=r"^#[0-9A-Fa-f]{6}$")] = "#173B65"
    secondary_color: Annotated[str, Field(pattern=r"^#[0-9A-Fa-f]{6}$")] = "#EAF1F8"
    footer_text: Annotated[str, Field(max_length=200)] = ""
    include_generated_date: bool = True

    @field_validator("brand_name", "logo_path", "footer_text")
    @classmethod
    def valid_xml_text(cls, value: str) -> str:
        if INVALID_XML_CONTROL.search(value):
            raise ValueError("展示参数包含不支持的控制字符。")
        return value

    @field_validator("image_paths")
    @classmethod
    def valid_image_indexes(cls, value: dict[int, str]) -> dict[int, str]:
        if any(index < 0 or index >= 80 for index in value):
            raise ValueError("示意图索引无效。")
        return value


class ActivityPlanDocumentRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")

    content: ActivityPlanContent
    presentation: DocumentPresentationOptions = Field(default_factory=DocumentPresentationOptions)


class DocumentGenerationError(ValueError):
    pass


class DocumentNotFound(LookupError):
    pass


@dataclass(frozen=True, slots=True)
class GeneratedDocument:
    document_id: str
    filename: str
    content_type: str
    download_url: str
    storage_key: str
    content: bytes | None = None

    def public(self) -> dict[str, str]:
        return {
            "document_id": self.document_id,
            "filename": self.filename,
            "content_type": self.content_type,
            "download_url": self.download_url,
        }


@dataclass(frozen=True, slots=True)
class ActivityPlanTemplateSelection:
    variant: str
    fallback_reason: str | None = None
    logo_bytes: bytes | None = None
    brand_name: str = ""


def _set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, *, top: int = 100, start: int = 110, bottom: int = 100, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = "D9D9D9", size: str = "6") -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_row_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _remove_paragraph_borders(paragraph_or_style) -> None:
    properties = paragraph_or_style._element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _append_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def _set_font(run, *, east_asia: str = "Noto Sans CJK SC", latin: str = "Arial", size: float | None = None,
              bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_text(paragraph, value: str, *, size: float = 11, color: str = "222222", bold: bool = False) -> None:
    parts = value.split("[待确认]")
    for index, part in enumerate(parts):
        if index:
            marker = paragraph.add_run("[待确认]")
            _set_font(marker, size=size, color="A23B20", bold=True)
        if part:
            run = paragraph.add_run(part)
            _set_font(run, size=size, color=color, bold=bold)


class ActivityPlanDocxRenderer:
    """Render validated content and presentation options without content generation."""

    def __init__(self, clock: Callable[[], datetime] | None = None) -> None:
        self._clock = clock or (lambda: datetime.now(timezone.utc))

    def render(
        self,
        content: ActivityPlanContent,
        presentation: DocumentPresentationOptions,
        *,
        logo_bytes: bytes | None = None,
    ) -> bytes:
        document = Document()
        self._configure_document(document, content, presentation)
        self._add_cover(document, content, presentation, logo_bytes)
        document.add_page_break()
        self._add_text_section(document, "活动主题", content.activity_theme)
        self._add_text_section(document, "活动时间", content.activity_time)
        self._add_text_section(document, "活动地点", content.activity_location)
        self._add_text_section(document, "活动对象", content.target_audience)
        self._add_list_section(document, "活动宣发途径", content.promotion_channels)
        self._add_execution_table(document, content.activity_items, presentation)
        self._add_text_section(document, "邀约文案", content.invitation_copy)
        if content.pending_items:
            self._add_list_section(document, "待确认事项", content.pending_items, pending=True)
        output = io.BytesIO()
        document.save(output)
        value = output.getvalue()
        if not value or len(value) > MAX_DOCUMENT_BYTES:
            raise DocumentGenerationError("生成的 DOCX 文件大小异常。")
        return value

    def _configure_document(
        self, document: Document, content: ActivityPlanContent, presentation: DocumentPresentationOptions
    ) -> None:
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(19)
        section.bottom_margin = Mm(18)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)

        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.space_after = Pt(6)

        title = document.styles["Title"]
        title.font.name = "Arial"
        title.font.size = Pt(26)
        title.font.bold = True
        title.font.color.rgb = RGBColor(0, 0, 0)
        title._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        _remove_paragraph_borders(title)

        heading = document.styles["Heading 1"]
        heading.font.name = "Arial"
        heading.font.size = Pt(15)
        heading.font.bold = True
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(6)
        heading.paragraph_format.keep_with_next = True

        properties = document.core_properties
        properties.title = content.document_title
        properties.subject = content.activity_theme
        properties.author = presentation.brand_name or "Enterprise AI Agent Workbench"
        properties.keywords = "activity plan, docx"

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_text = presentation.brand_name or "企业活动方案"
        if content.document_title and content.document_title != header_text:
            header_text = f"{header_text}  |  {content.document_title}"
        run = header.add_run(header_text)
        _set_font(run, size=8.5, color="000000")

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if presentation.footer_text:
            prefix = footer.add_run(f"{presentation.footer_text}  ·  ")
            _set_font(prefix, size=8.5, color="666666")
        page_prefix = footer.add_run("第 ")
        _set_font(page_prefix, size=8.5, color="666666")
        _append_page_field(footer)
        page_suffix = footer.add_run(" 页")
        _set_font(page_suffix, size=8.5, color="666666")

    def _add_cover(
        self,
        document: Document,
        content: ActivityPlanContent,
        presentation: DocumentPresentationOptions,
        logo_bytes: bytes | None,
    ) -> None:
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(28)
        if logo_bytes:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image = Image.from_blob(logo_bytes)
            if image.px_width > 20_000 or image.px_height > 20_000:
                raise DocumentGenerationError("Logo 像素尺寸超过限制。")
            max_width, max_height = Mm(38), Mm(22)
            factor = min(float(max_width) / image.width, float(max_height) / image.height, 1.0)
            paragraph.add_run().add_picture(
                io.BytesIO(logo_bytes), width=Emu(int(image.width * factor)), height=Emu(int(image.height * factor))
            )
            paragraph.paragraph_format.space_after = Pt(24)

        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _remove_paragraph_borders(title)
        _add_text(title, content.document_title, size=26, color="000000", bold=True)
        title.paragraph_format.space_after = Pt(20)

        theme = document.add_paragraph()
        theme.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_text(theme, content.activity_theme, size=16, color="000000", bold=True)
        theme.paragraph_format.space_after = Pt(18)

        if presentation.brand_name:
            brand = document.add_paragraph()
            brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_text(brand, presentation.brand_name, size=11, color="555555")
            brand.paragraph_format.space_after = Pt(8)
        if presentation.include_generated_date:
            generated = document.add_paragraph()
            generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
            generated_at = self._clock().astimezone(timezone.utc).date().isoformat()
            _add_text(generated, f"生成日期 {generated_at}", size=9.5, color="777777")

    @staticmethod
    def _heading(document: Document, title: str):
        paragraph = document.add_paragraph(title, style="Heading 1")
        paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def _add_text_section(self, document: Document, title: str, value: str) -> None:
        self._heading(document, title)
        paragraph = document.add_paragraph()
        _add_text(paragraph, value)
        paragraph.paragraph_format.keep_together = True

    def _add_list_section(self, document: Document, title: str, values: list[str], *, pending: bool = False) -> None:
        self._heading(document, title)
        if not values:
            document.add_paragraph()
            return
        keep_pending_block = pending and len(values) <= 12 and sum(len(value) for value in values) <= 2_000
        for index, value in enumerate(values):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_text(paragraph, value, color="222222" if not pending else "6B2F20")
            paragraph.paragraph_format.space_after = Pt(4)
            if keep_pending_block:
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = index < len(values) - 1

    def _add_execution_table(
        self, document: Document, items: list[ActivityPlanItem], presentation: DocumentPresentationOptions
    ) -> None:
        self._heading(document, "活动执行表")
        table = document.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = (Mm(25), Mm(35), Mm(70), Mm(40))
        headers = ("环节", "名称", "说明", "示意图需求")
        primary = presentation.primary_color.lstrip("#").upper()
        secondary = presentation.secondary_color.lstrip("#").upper()
        _set_table_borders(table)
        header_row = table.rows[0]
        _set_repeat_table_header(header_row)
        for index, cell in enumerate(header_row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_shading(cell, primary)
            _set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(headers[index])
            _set_font(run, size=10, color="FFFFFF", bold=True)

        for row_index, item in enumerate(items):
            values = (item.phase, item.name, item.description, item.image_requirement)
            row = table.add_row()
            if sum(len(value) for value in values) <= 1_800:
                _set_row_cant_split(row)
            cells = row.cells
            for index, (cell, value) in enumerate(zip(cells, values)):
                cell.width = widths[index]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_margins(cell)
                if row_index % 2:
                    _set_cell_shading(cell, secondary)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in {0, 1} else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
                _add_text(paragraph, value, size=9.5)


def sanitize_filename_component(value: str, *, max_length: int = 60) -> str:
    cleaned = WINDOWS_ILLEGAL_FILENAME.sub("_", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._")
    if not cleaned:
        cleaned = "未命名活动"
    if cleaned.upper() in WINDOWS_RESERVED_NAMES:
        cleaned = f"_{cleaned}"
    return cleaned[:max_length].rstrip(" .") or "未命名活动"


def validate_logo_storage_key(value: str, tenant_id: str) -> str:
    if not value:
        return ""
    if "\\" in value or len(value) > 512:
        raise DocumentGenerationError("Logo path 无效。")
    path = PurePosixPath(value)
    parts = path.parts
    if path.is_absolute() or any(part in {"", ".", ".."} for part in parts):
        raise DocumentGenerationError("Logo path 无效。")
    expected = ("brand-assets", tenant_id, "logos")
    if len(parts) != 4 or tuple(parts[:3]) != expected or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
        raise DocumentGenerationError("Logo 只能引用当前企业 brand-assets 下的 PNG/JPEG 对象。")
    return path.as_posix()


def validate_illustration_storage_key(value: str, tenant_id: str) -> str:
    if not isinstance(value, str) or "\\" in value or len(value) > 512:
        raise DocumentGenerationError("示意图路径无效。")
    path = PurePosixPath(value)
    parts = path.parts
    expected = ("brand-assets", tenant_id, "activity-plan-images")
    if (path.is_absolute() or any(part in {"", ".", ".."} for part in parts)
            or len(parts) != 4 or tuple(parts[:3]) != expected
            or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}):
        raise DocumentGenerationError("示意图只能引用当前企业的活动图片对象。")
    return path.as_posix()


class ActivityPlanDocumentService:
    def __init__(
        self,
        storage: StorageProvider,
        *,
        renderer: ActivityPlanDocxRenderer | None = None,
        clock: Callable[[], datetime] | None = None,
        brand_provider: Callable[[str], dict] | None = None,
    ) -> None:
        self._storage = storage
        self._clock = clock or (lambda: datetime.now(timezone.utc))
        self._renderer = renderer or ActivityPlanDocxRenderer(self._clock)
        self._brand_provider = brand_provider

    def resolve_activity_plan_template(self, tenant_id: str) -> ActivityPlanTemplateSelection:
        """Select from the current tenant asset, without trusting request presentation paths."""
        if self._brand_provider is None:
            return ActivityPlanTemplateSelection("legacy", "logo_missing")
        try:
            config = self._brand_provider(tenant_id)
        except Exception:
            return ActivityPlanTemplateSelection("legacy", "logo_unreadable")
        if not isinstance(config, dict):
            return ActivityPlanTemplateSelection("legacy", "logo_invalid")
        key = config.get("brand_logo")
        if not key:
            return ActivityPlanTemplateSelection("legacy", "logo_missing")
        try:
            key = validate_brand_logo_key(key, tenant_id)
        except BrandLogoError:
            return ActivityPlanTemplateSelection("legacy", "ownership_invalid")
        try:
            logo_bytes = self._storage.get(key)
        except Exception:
            return ActivityPlanTemplateSelection("legacy", "logo_unreadable")
        try:
            inspect_transparent_png(logo_bytes)
        except BrandLogoError:
            return ActivityPlanTemplateSelection("legacy", "logo_invalid")
        brand_name = config.get("brand_name")
        if (not isinstance(brand_name, str) or not brand_name.strip() or len(brand_name) > 120
                or INVALID_XML_CONTROL.search(brand_name)):
            return ActivityPlanTemplateSelection("legacy", "logo_invalid")
        return ActivityPlanTemplateSelection("branded", logo_bytes=logo_bytes, brand_name=brand_name.strip())

    def _illustrations(self, tenant_id: str, content: ActivityPlanContent,
                       presentation: DocumentPresentationOptions) -> dict[int, bytes]:
        result = {}
        for index, path in presentation.image_paths.items():
            if index >= len(content.activity_items):
                raise DocumentGenerationError("示意图索引超出活动环节范围。")
            key = validate_illustration_storage_key(path, tenant_id)
            picture = self._storage.get(key)
            if not picture or len(picture) > MAX_LOGO_BYTES:
                raise DocumentGenerationError("示意图文件为空或超过 2MB。")
            result[index] = picture
        return result

    def create(
        self,
        tenant_id: str,
        user_id: str,
        content: ActivityPlanContent,
        presentation: DocumentPresentationOptions,
    ) -> GeneratedDocument:
        selection = self.resolve_activity_plan_template(tenant_id)
        TEMPLATE_LOGGER.info("activity_plan_template_selection %s", json.dumps({
            "tenant_id": tenant_id,
            "template_variant": selection.variant,
            "fallback_reason": selection.fallback_reason,
        }, ensure_ascii=False, sort_keys=True))
        try:
            if selection.variant == "branded":
                from app.branded_activity_plan_template import BrandedActivityPlanDocxRenderer

                branded_presentation = presentation.model_copy(update={
                    "brand_name": selection.brand_name,
                    "logo_path": "",
                })
                value = BrandedActivityPlanDocxRenderer(self._clock).render(
                    content, branded_presentation, logo_bytes=selection.logo_bytes,
                    illustration_bytes=self._illustrations(tenant_id, content, presentation),
                )
            else:
                # Preserve the Production renderer and its existing request behavior.
                logo_bytes = None
                logo_key = validate_logo_storage_key(presentation.logo_path, tenant_id)
                if logo_key:
                    logo_bytes = self._storage.get(logo_key)
                    if not logo_bytes or len(logo_bytes) > MAX_LOGO_BYTES:
                        raise DocumentGenerationError("Logo 文件为空或超过 2MB。")
                value = self._renderer.render(content, presentation, logo_bytes=logo_bytes)
        except DocumentGenerationError:
            raise
        except (StorageObjectNotFound, StorageUnavailable):
            raise
        except Exception as exc:
            raise DocumentGenerationError("DOCX 生成失败。") from exc
        timestamp = self._clock().astimezone(timezone.utc).strftime("%Y%m%d-%H%M%S")
        theme = sanitize_filename_component(content.activity_theme or content.document_title)
        filename = f"活动方案_{theme}_{timestamp}.docx"
        document_id = uuid4().hex
        storage_key = f"generated-documents/{tenant_id}/{user_id}/{document_id}/{filename}"
        self._storage.put(storage_key, value, DOCX_CONTENT_TYPE)
        return GeneratedDocument(
            document_id=document_id,
            filename=filename,
            content_type=DOCX_CONTENT_TYPE,
            download_url=f"/api/v1/documents/activity-plan/{document_id}",
            storage_key=storage_key,
        )

    def get(self, tenant_id: str, user_id: str, document_id: str) -> GeneratedDocument:
        if not DOCUMENT_ID_RE.fullmatch(document_id):
            raise DocumentNotFound("文档不存在。")
        prefix = f"generated-documents/{tenant_id}/{user_id}/{document_id}/"
        keys = [key for key in self._storage.list_keys(prefix) if key.startswith(prefix) and key.endswith(".docx")]
        if len(keys) != 1:
            raise DocumentNotFound("文档不存在。")
        key = keys[0]
        filename = PurePosixPath(key).name
        if sanitize_filename_component(filename.removesuffix(".docx"), max_length=180) + ".docx" != filename:
            raise DocumentNotFound("文档不存在。")
        try:
            content = self._storage.get(key)
        except StorageObjectNotFound as exc:
            raise DocumentNotFound("文档不存在。") from exc
        return GeneratedDocument(
            document_id=document_id,
            filename=filename,
            content_type=DOCX_CONTENT_TYPE,
            download_url=f"/api/v1/documents/activity-plan/{document_id}",
            storage_key=key,
            content=content,
        )
