"""Reference-style branded Activity Plan DOCX renderer.

The legacy renderer and structured content contract remain in document_generator.
This module is selected only after the current tenant logo passes capability checks.
"""
from __future__ import annotations

import io
from copy import deepcopy
from datetime import datetime, timezone
from typing import Callable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor
from PIL import Image as PillowImage

from app.document_generator import (
    ActivityPlanContent,
    ActivityPlanItem,
    DocumentGenerationError,
    DocumentPresentationOptions,
    MAX_DOCUMENT_BYTES,
)


def _set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, *, top: int = 45, start: int = 108, bottom: int = 45, end: int = 108) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = "000000", size: str = "4") -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_row_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _remove_paragraph_borders(paragraph_or_style) -> None:
    properties = paragraph_or_style._element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _append_page_field(paragraph):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._r.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    paragraph.add_run()._r.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    result = paragraph.add_run("1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)
    return result


def _set_font(run, *, east_asia: str = "FangSong", latin: str = "FangSong", size: float | None = None,
              bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_text(paragraph, value: str, *, size: float = 12, color: str = "000000", bold: bool = False) -> None:
    run = paragraph.add_run(value)
    _set_font(run, size=size, color=color, bold=bold)


def _line_format(paragraph) -> None:
    paragraph.paragraph_format.line_spacing = Pt(18.7)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def _fit_size(image_bytes: bytes, max_width: Emu, max_height: Emu) -> tuple[Emu, Emu]:
    image = Image.from_blob(image_bytes)
    if image.px_width > 20_000 or image.px_height > 20_000:
        raise DocumentGenerationError("图片像素尺寸超过限制。")
    factor = min(float(max_width) / image.width, float(max_height) / image.height)
    return Emu(int(image.width * factor)), Emu(int(image.height * factor))


def _opacity_png(image_bytes: bytes, opacity: float) -> bytes:
    """Normalize existing alpha, then apply template opacity without stretching."""
    try:
        with PillowImage.open(io.BytesIO(image_bytes)) as source:
            image = source.convert("RGBA")
        alpha = image.getchannel("A")
        peak = alpha.getextrema()[1]
        if not peak:
            raise DocumentGenerationError("Brand Logo 不包含可见像素。")
        alpha = alpha.point(lambda value: min(255, round(value * 255 * opacity / peak)))
        image.putalpha(alpha)
        output = io.BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()
    except DocumentGenerationError:
        raise
    except Exception as exc:
        raise DocumentGenerationError("Brand Logo 格式无效。") from exc


def _anchor_picture(shape, *, x: Emu, y: Emu, behind_text: bool) -> None:
    """Turn a header picture into a page-relative repeated anchor."""
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0",
        "relativeHeight": "0", "behindDoc": "1" if behind_text else "0",
        "locked": "1", "layoutInCell": "0", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis, offset in (("H", x), ("V", y)):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        value = OxmlElement("wp:posOffset")
        value.text = str(int(offset))
        position.append(value)
        anchor.append(position)
    anchor.append(deepcopy(inline.extent))
    anchor.append(OxmlElement("wp:wrapNone"))
    for child in inline:
        if child.tag in {qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")}:
            anchor.append(deepcopy(child))
    inline.getparent().replace(inline, anchor)


class BrandedActivityPlanDocxRenderer:
    """Render the two-page reference layout without changing validated content."""

    def __init__(self, clock: Callable[[], datetime] | None = None) -> None:
        self._clock = clock or (lambda: datetime.now(timezone.utc))

    def render(
        self,
        content: ActivityPlanContent,
        presentation: DocumentPresentationOptions,
        *,
        logo_bytes: bytes | None = None,
        watermark_bytes: bytes | None = None,
        illustration_bytes: dict[int, bytes] | None = None,
    ) -> bytes:
        if not logo_bytes:
            raise DocumentGenerationError("BRANDED_TEMPLATE_LOGO_UNAVAILABLE")
        if not presentation.brand_name.strip():
            raise DocumentGenerationError("BRAND_NAME_REQUIRED_FOR_ACTIVITY_PLAN_TEMPLATE")
        document = Document()
        self._configure_document(document, content, presentation, logo_bytes, watermark_bytes or logo_bytes)
        title = document.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _remove_paragraph_borders(title)
        title.paragraph_format.line_spacing = Pt(26)
        title.paragraph_format.space_before = Pt(2.7)
        title.paragraph_format.space_after = Pt(20)
        _add_text(title, content.document_title, size=20.05, bold=True)

        self._numbered(document, "一、活动主题：")
        theme = document.add_paragraph()
        _line_format(theme)
        theme.paragraph_format.left_indent = Mm(12.7)
        _add_text(theme, content.activity_theme or "【待确认】")
        self._numbered(document, "二、活动时间：", content.activity_time)
        self._numbered(document, "三、活动地点：", content.activity_location)
        self._numbered(document, "四、活动对象：", content.target_audience)
        self._numbered(document, "五、活动宣发途径")
        for channel in content.promotion_channels:
            bullet = document.add_paragraph()
            _line_format(bullet)
            bullet.paragraph_format.left_indent = Mm(16.25)
            bullet.paragraph_format.first_line_indent = Mm(-7.8)
            bullet.paragraph_format.tab_stops.add_tab_stop(Mm(16.25))
            # Unicode triangle remains readable when Wingdings is unavailable.
            glyph = bullet.add_run("▸")
            _set_font(glyph, east_asia="Segoe UI Symbol", latin="Segoe UI Symbol", size=12)
            bullet.add_run("\t")
            _add_text(bullet, channel)
        self._numbered(document, "六、活动说明")
        self._add_execution_table(document, content.activity_items, illustration_bytes or {})
        self._add_attachment(document, content.invitation_copy)
        output = io.BytesIO()
        document.save(output)
        value = output.getvalue()
        if not value or len(value) > MAX_DOCUMENT_BYTES:
            raise DocumentGenerationError("生成的 DOCX 文件大小异常。")
        return value

    def _configure_document(
        self, document: Document, content: ActivityPlanContent, presentation: DocumentPresentationOptions,
        logo_bytes: bytes, watermark_bytes: bytes,
    ) -> None:
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.left_margin = Mm(31.75)
        section.right_margin = Mm(31.75)
        section.header_distance = Mm(13.8)
        section.footer_distance = Mm(15.95)

        normal = document.styles["Normal"]
        normal.font.name = "FangSong"
        normal.font.size = Pt(12)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "FangSong")
        normal.paragraph_format.line_spacing = Pt(18.7)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)

        title = document.styles["Title"]
        title.font.name = "FangSong"
        title.font.size = Pt(20.05)
        title.font.bold = True
        title.font.color.rgb = RGBColor(0, 0, 0)
        title._element.rPr.rFonts.set(qn("w:eastAsia"), "FangSong")
        _remove_paragraph_borders(title)

        properties = document.core_properties
        properties.title = content.document_title
        properties.subject = content.activity_theme
        properties.author = presentation.brand_name or "Enterprise AI Agent Workbench"
        properties.keywords = "activity plan, docx"

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _line_format(header)
        _add_text(header, f"{presentation.brand_name}-活动方案  请勿外传", size=9, color="7F7F7F")
        visible_logo = _opacity_png(logo_bytes, 0.60)
        logo_width, logo_height = _fit_size(visible_logo, Mm(50.9), Mm(11.7))
        logo = header.add_run().add_picture(io.BytesIO(visible_logo), width=logo_width, height=logo_height)
        _anchor_picture(logo, x=Mm(0.8), y=Mm(1.8), behind_text=False)
        # A normal enterprise logo is faint at 10% opacity. The supplied
        # WPS reference uses a separate pastel mark whose embedded alpha peaks
        # at 50/255; preserve that mark's appearance when it is supplied.
        watermark_opacity = 0.20 if watermark_bytes != logo_bytes else 0.10
        faint_logo = _opacity_png(watermark_bytes, watermark_opacity)
        mark_width, mark_height = _fit_size(faint_logo, Mm(146.5), Mm(146.5))
        mark = header.add_run().add_picture(io.BytesIO(faint_logo), width=mark_width, height=mark_height)
        _anchor_picture(mark, x=Mm(31.75), y=Mm(75.3), behind_text=True)

        footer = section.footer.paragraphs[0]
        footer_style = document.styles["Footer"]
        footer_style.font.name = "Calibri"
        footer_style.font.size = Pt(9)
        footer_style.font.color.rgb = RGBColor(0, 0, 0)
        footer_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _line_format(footer)
        page = _append_page_field(footer)
        _set_font(page, east_asia="Calibri", latin="Calibri", size=9, color="000000")

    @staticmethod
    def _numbered(document: Document, label: str, value: str | None = None) -> None:
        paragraph = document.add_paragraph()
        _line_format(paragraph)
        paragraph.paragraph_format.keep_with_next = not value
        _add_text(paragraph, label, bold=True)
        if value is not None:
            _add_text(paragraph, value or "【待确认】")

    @staticmethod
    def _add_attachment(document: Document, invitation_copy: str) -> None:
        heading = document.add_paragraph()
        heading.paragraph_format.page_break_before = True
        heading.paragraph_format.keep_with_next = True
        _line_format(heading)
        _add_text(heading, "附件一：物业邀约文案（参考）", bold=True)
        for line in invitation_copy.splitlines() or [invitation_copy]:
            paragraph = document.add_paragraph()
            _line_format(paragraph)
            paragraph.paragraph_format.line_spacing = Pt(17)
            _add_text(paragraph, line)

    def _add_execution_table(
        self, document: Document, items: list[ActivityPlanItem], illustration_bytes: dict[int, bytes]
    ) -> None:
        table = document.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), str(round(Mm(2.05) / 635)))
        indent.set(qn("w:type"), "dxa")
        table._tbl.tblPr.append(indent)
        widths = (Mm(15.8), Mm(22.9), Mm(78.7), Mm(40.2))
        for index, width in enumerate(widths):
            table.columns[index].width = width
        headers = ("环节", "名称", "说明", "示意图")
        _set_table_borders(table)
        header_row = table.rows[0]
        header_row.height = Mm(6.8)
        header_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        _set_repeat_table_header(header_row)
        for index, cell in enumerate(header_row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_shading(cell, "D8D8D8")
            _set_cell_margins(cell, top=0, bottom=0)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _line_format(paragraph)
            run = paragraph.add_run(headers[index])
            _set_font(run, size=12, color="000000", bold=True)

        for row_index, item in enumerate(items):
            row = table.add_row()
            if len(item.description) <= 1_800:
                _set_row_cant_split(row)
            for index, width in enumerate(widths):
                row.cells[index].width = width
            name_cell = row.cells[0].merge(row.cells[1])
            description_cell = row.cells[2]
            image_cell = row.cells[3]
            for index, (cell, value) in enumerate(((name_cell, item.name), (description_cell, item.description))):
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_margins(cell)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                _line_format(paragraph)
                if index == 1 and "\n" in value:
                    paragraph.paragraph_format.space_after = Pt(12.5)
                _add_text(paragraph, value)
            image_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins(image_cell)
            image_paragraph = image_cell.paragraphs[0]
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _line_format(image_paragraph)
            picture = illustration_bytes.get(row_index)
            if picture:
                image_paragraph.paragraph_format.line_spacing = Pt(18.7)
                image_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                source = Image.from_blob(picture)
                ratio = source.px_height / source.px_width
                max_height = Mm(36) if ratio > 1.29 else Mm(41.5)
                width, height = _fit_size(picture, Mm(35.5), max_height)
                if width > Pt(90):
                    # LibreOffice adds an extra cell inset for wide inline
                    # pictures; counter it without changing the image ratio.
                    image_paragraph.paragraph_format.left_indent = Mm(-2.5)
                image_paragraph.add_run().add_picture(io.BytesIO(picture), width=width, height=height)
            else:
                _add_text(image_paragraph, "/")
