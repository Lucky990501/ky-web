"""Tenant-scoped enterprise document ingestion and retrieval services."""
from __future__ import annotations

import asyncio
import hashlib
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path

from app.knowledge_metadata import (
    RAG_INDEX_VERSION,
    build_chunk_metadata,
    detect_query_intent,
    embedding_input,
    keyword_score as structured_keyword_score,
    metadata_dict,
    metadata_score as calculate_metadata_score,
    query_embedding_input,
)
from app.product_store import ProductStore
from app.settings import Settings, safe_runtime_config_snapshot
from app.storage import storage_provider
import httpx


_embedding_probe_error: str | None = None


def set_embedding_probe(error: str | None) -> None:
    """Keep only a safe availability state, never provider response details."""
    global _embedding_probe_error
    _embedding_probe_error = error


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    title: str
    sections: tuple[dict, ...]


class DocumentParser:
    """Format router; parsing remains outside HTTP controllers."""
    async def parse(self, filename: str, content: bytes) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig", errors="replace")
            return ParsedDocument(normalize_text(text), Path(filename).stem, tuple(sectionize(text)))
        if suffix == ".pdf":
            return await asyncio.to_thread(self._parse_pdf, filename, content)
        if suffix == ".docx":
            return await asyncio.to_thread(self._parse_docx, filename, content)
        raise ValueError("仅支持 PDF、DOCX、TXT、MD 文件。")

    @staticmethod
    def _parse_pdf(filename: str, content: bytes) -> ParsedDocument:
        try:
            from pypdf import PdfReader
            from io import BytesIO
            reader = PdfReader(BytesIO(content))
            pages = [{"text": page.extract_text() or "", "page_number": index + 1} for index, page in enumerate(reader.pages)]
        except Exception as exc:
            raise ValueError("PDF 解析失败，请确认文件未损坏且包含可提取文本。") from exc
        text = "\n\n".join(item["text"] for item in pages)
        return ParsedDocument(normalize_text(text), Path(filename).stem, tuple(pages))

    @staticmethod
    def _parse_docx(filename: str, content: bytes) -> ParsedDocument:
        try:
            from docx import Document
            from io import BytesIO
            document = Document(BytesIO(content))
            lines = [item.text for item in document.paragraphs if item.text.strip()]
            for table in document.tables:
                lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        except Exception as exc:
            raise ValueError("DOCX 解析失败，请确认文件格式正确。") from exc
        text = "\n".join(lines)
        return ParsedDocument(normalize_text(text), Path(filename).stem, tuple(sectionize(text)))


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sectionize(text: str) -> list[dict]:
    title = "正文"
    sheet_name: str | None = None
    parts: list[dict] = []
    buffer: list[str] = []
    for line in normalize_text(text).splitlines():
        stripped = line.strip()
        markdown_heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        numbered_heading = re.match(r"^[一二三四五六七八九十0-9]+[、.]\s*(.+)$", stripped)
        if markdown_heading or numbered_heading:
            if buffer:
                parts.append({"text": "\n".join(buffer), "section": title, "sheet_name": sheet_name})
            if markdown_heading:
                level, heading = len(markdown_heading.group(1)), markdown_heading.group(2).strip()
                if level == 1:
                    sheet_name = heading
                title = heading
            else:
                title = numbered_heading.group(1).strip()
            buffer = []
        else:
            buffer.append(line)
    if buffer:
        parts.append({"text": "\n".join(buffer), "section": title, "sheet_name": sheet_name})
    return parts or [{"text": normalize_text(text), "section": "正文", "sheet_name": sheet_name}]


class ChunkingStrategy:
    def __init__(self, size: int, overlap: int) -> None:
        self.size, self.overlap = size, overlap

    def split(self, parsed: ParsedDocument) -> list[dict]:
        chunks: list[dict] = []
        for source in parsed.sections:
            text = normalize_text(source.get("text", ""))
            while text:
                part = text[:self.size]
                if len(text) > self.size:
                    pivot = max(part.rfind("\n"), part.rfind("。"), part.rfind("；"))
                    if pivot > self.size // 2:
                        part = text[:pivot + 1]
                chunks.append({"content": part, "title": parsed.title, "section": source.get("section", "正文"), "page_number": source.get("page_number"), "chunk_index": len(chunks), "metadata": {"source": "enterprise_file", "sheet_name": source.get("sheet_name")}})
                if len(part) >= len(text):
                    break
                text = text[max(1, len(part) - self.overlap):].lstrip()
        return chunks


class EmbeddingProvider:
    async def embed_documents(self, texts: list[str]) -> list[list[float]]: raise NotImplementedError
    async def embed_query(self, text: str) -> list[float]: raise NotImplementedError


@dataclass(frozen=True, slots=True)
class RetrievalConfidencePolicy:
    """Configurable V1 gate that rejects weak, ungrounded retrieval candidates."""

    minimum_final_score: float
    minimum_vector_score: float
    keyword_exact_match: bool
    result_margin: float

    @classmethod
    def from_settings(cls, settings: Settings) -> "RetrievalConfidencePolicy":
        return cls(
            minimum_final_score=settings.knowledge_min_final_score,
            minimum_vector_score=settings.knowledge_min_vector_score,
            keyword_exact_match=settings.knowledge_keyword_exact_match,
            result_margin=settings.knowledge_result_margin,
        )

    def decision(self, *, vector_score: float, keyword_score: float, final_score: float) -> tuple[bool, str | None]:
        if final_score < self.minimum_final_score:
            return False, "below_minimum_final_score"
        if vector_score < self.minimum_vector_score and keyword_score <= 0:
            return False, "below_minimum_vector_score"
        if self.keyword_exact_match and keyword_score <= 0:
            return False, "keyword_exact_match_required"
        if keyword_score <= 0 and final_score < self.minimum_final_score + self.result_margin:
            return False, "within_confidence_margin"
        return True, None


@dataclass(frozen=True, slots=True)
class RetrievalOutcome:
    """A retrieval decision safe to expose to internal evaluation tooling."""

    results: list[dict]
    accepted: bool
    rejection_reason: str | None = None
    rejection_detail: str | None = None


@dataclass(frozen=True, slots=True)
class QueryAnswerabilityPolicy:
    """Reject unsupported claims before ungrounded context reaches an Agent."""

    VERSION = "query-guard-v2"
    enabled: bool = True

    _absolute_promise = re.compile(r"(?:保证|承诺|包|保).{0,12}(?:提升|提分|保过|录取|收益|前[0-9一二三四五六七八九十]+)")
    _explicitly_nonexistent = re.compile(r"不存在.{0,12}(?:课程|老师|教师|讲师|活动|项目|名额|班级)")
    _private_or_credential_data = re.compile(r"(?:私人|个人|他人|用户|客户|员工|负责人).{0,16}(?:电话|手机号|联系方式|住址|身份证|银行卡|密码|验证码|名单)|(?:密码|验证码|密钥|api[ _-]?key|token|银行卡号|cvv)")
    _non_public_enterprise_data = re.compile(r"(?:未公开|内部|机密|保密|竞争对手).{0,20}(?:融资|金额|客户|名单|数据|资料|信息|财务|报价)")
    _obvious_out_of_scope = re.compile(r"(?:预订|购买|维修|报修|转账|开户).{0,20}(?:机票|航班|酒店|设备|金融|银行卡)")
    _price_request = re.compile(r"(?:价格|费用|收费|报价|多少钱|人民币|[0-9]+\s*元)")
    _price_evidence = re.compile(r"(?:价格|费用|收费|报价|人民币|[0-9]+\s*元)")

    def pre_retrieval_rejection_reason(self, query: str) -> str | None:
        """Reject category-level unsafe or unanswerable requests before embedding."""
        if not self.enabled:
            return None
        normalized = re.sub(r"\s+", "", query)
        if self._non_public_enterprise_data.search(normalized):
            return "non_public_enterprise_data"
        if self._private_or_credential_data.search(normalized):
            return "private_or_credential_data"
        if self._obvious_out_of_scope.search(normalized):
            return "obvious_out_of_scope"
        if self._explicitly_nonexistent.search(normalized):
            return "explicitly_nonexistent_entity"
        if self._absolute_promise.search(normalized):
            return "unsupported_absolute_promise"
        return None

    def rejection_reason(self, query: str, candidates: list[dict]) -> str | None:
        pre_retrieval_reason = self.pre_retrieval_rejection_reason(query)
        if pre_retrieval_reason:
            return pre_retrieval_reason
        if not self.enabled:
            return None
        normalized = re.sub(r"\s+", "", query)
        if self._price_request.search(normalized):
            evidence = " ".join(
                f"{item.get('title', '')} {item.get('section', '')} {item.get('content', '')}"
                for item in candidates
            )
            if not self._price_evidence.search(evidence):
                return "price_without_grounding"
        return None


def retrieval_policy_diagnostic(settings: Settings) -> dict:
    """Return the P0 RAG policy state without exposing secrets or raw URLs."""
    return {
        "query_guard_enabled": settings.knowledge_query_guard_enabled,
        "guard_policy_version": QueryAnswerabilityPolicy.VERSION,
        "min_final_score": settings.knowledge_min_final_score,
        "min_vector_score": settings.knowledge_min_vector_score,
        "result_margin": settings.knowledge_result_margin,
        "runtime_config_fingerprint": safe_runtime_config_snapshot()["fingerprint"],
    }


class LocalHashEmbeddingProvider(EmbeddingProvider):
    """Deterministic V1 provider; swappable through the provider interface."""
    def __init__(self, dimension: int) -> None: self.dimension = dimension
    def _embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimension
        for gram in re.findall(r"[\u4e00-\u9fff]{1,2}|[a-zA-Z0-9_]+", text.lower()):
            vector[int(hashlib.sha256(gram.encode()).hexdigest(), 16) % self.dimension] += 1
        norm = math.sqrt(sum(item * item for item in vector)) or 1.0
        return [round(item / norm, 8) for item in vector]
    async def embed_documents(self, texts: list[str]) -> list[list[float]]: return [self._embed(item) for item in texts]
    async def embed_query(self, text: str) -> list[float]: return self._embed(text)


class OpenAICompatibleEmbeddingProvider(EmbeddingProvider):
    """OpenAI-compatible `/embeddings` adapter; credentials stay in Settings."""

    def __init__(self, settings: Settings) -> None:
        self._base_url = settings.embedding_base_url.rstrip("/")
        self._api_key = settings.embedding_api_key
        self._model = settings.embedding_model
        self._dimension = settings.embedding_dimension

    def _payload(self, texts: list[str]) -> dict:
        return {"model": self._model, "input": texts, "encoding_format": "float"}

    def _vectors(self, body: dict, expected: int) -> list[list[float]]:
        data = sorted(body.get("data") or [], key=lambda item: item.get("index", 0))
        vectors = [item.get("embedding") for item in data]
        if len(vectors) != expected or any(not isinstance(vector, list) or len(vector) != self._dimension for vector in vectors):
            raise RuntimeError("Embedding Provider 返回的向量数量或维度与配置不一致。")
        return [[float(value) for value in vector] for vector in vectors]

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        async with httpx.AsyncClient(timeout=45) as client:
            response = await client.post(f"{self._base_url}/embeddings", headers={"Authorization": f"Bearer {self._api_key}"}, json=self._payload(texts))
            response.raise_for_status()
            return self._vectors(response.json(), len(texts))

    async def embed_query(self, text: str) -> list[float]:
        return (await self.embed_documents([text]))[0]

    def embed_query_sync(self, text: str) -> list[float]:
        with httpx.Client(timeout=45) as client:
            response = client.post(f"{self._base_url}/embeddings", headers={"Authorization": f"Bearer {self._api_key}"}, json=self._payload([text]))
            response.raise_for_status()
            return self._vectors(response.json(), 1)[0]


def embedding_provider_for(settings: Settings) -> EmbeddingProvider:
    if settings.embedding_provider == "local-hash":
        return LocalHashEmbeddingProvider(settings.embedding_dimension)
    if settings.embedding_provider in {"openai-compatible", "openai"}:
        return OpenAICompatibleEmbeddingProvider(settings)
    raise RuntimeError("未支持的 Embedding Provider。")


def runtime_diagnostic(product: ProductStore, settings: Settings) -> dict:
    """Describe readiness without leaking credentials or silently downgrading."""
    reasons: list[str] = []
    pgvector_installed = False
    if not product._store.is_postgres:
        reasons.append("生产语义检索需要 PostgreSQL 和 pgvector。")
    else:
        try:
            with product._store.connection() as conn:
                pgvector_installed = bool(conn.execute("SELECT EXISTS(SELECT 1 FROM pg_extension WHERE extname='vector') AS installed").fetchone()["installed"])
        except Exception:
            reasons.append("无法确认 pgvector 扩展状态。")
        if not pgvector_installed:
            reasons.append("PostgreSQL 未启用 pgvector 扩展。")
    if settings.embedding_provider == "local-hash":
        reasons.append("当前 Embedding Provider 为 local-hash，不是正式 Embedding Provider。")
    if settings.embedding_provider != "local-hash" and not settings.embedding_api_key:
        reasons.append("正式 Embedding Provider 未配置 EMBEDDING_API_KEY。")
    if settings.embedding_provider != "local-hash" and not settings.embedding_base_url:
        reasons.append("正式 Embedding Provider 未配置 EMBEDDING_BASE_URL。")
    if _embedding_probe_error:
        reasons.append("正式 Embedding Provider 连通性验证失败。")
    strict = settings.environment == "production" and not settings.knowledge_allow_fallback
    return {
        "status": "ok" if not reasons else "degraded",
        "strict": strict,
        "pgvector_installed": pgvector_installed,
        "embedding_provider": settings.embedding_provider,
        "embedding_model": settings.embedding_model,
        "embedding_dimension": settings.embedding_dimension,
        "allow_fallback": settings.knowledge_allow_fallback,
        "reasons": reasons,
    }


def require_semantic_runtime(product: ProductStore, settings: Settings) -> None:
    diagnostic = runtime_diagnostic(product, settings)
    if diagnostic["strict"] and diagnostic["status"] != "ok":
        raise RuntimeError("生产语义检索未就绪：" + "；".join(diagnostic["reasons"]))


class KnowledgeProcessingService:
    EMBEDDING_BATCH_SIZE = 32

    def __init__(self, product: ProductStore, settings: Settings) -> None:
        self.product, self.settings = product, settings
        self.parser = DocumentParser(); self.chunker = ChunkingStrategy(settings.knowledge_chunk_size, settings.knowledge_chunk_overlap)
        self.embedding = embedding_provider_for(settings)

    async def process(self, tenant_id: str, file_id: str, *, force: bool = False, raise_errors: bool = False) -> None:
        record = self.product.knowledge_file(tenant_id, file_id)
        if not record or (record["status"] == "ready" and not force): return
        previous_status = record["status"]
        try:
            require_semantic_runtime(self.product, self.settings)
            if not force:
                self.product.set_knowledge_file_status(tenant_id, file_id, "parsing")
            content = storage_provider(self.settings).get(record["storage_key"])
            parsed = await self.parser.parse(record.get("filename") or record["name"], content)
            if not force:
                self.product.set_knowledge_file_status(tenant_id, file_id, "chunking", parsed_text=parsed.text)
            chunks = self.chunker.split(parsed)
            if not force:
                self.product.set_knowledge_file_status(tenant_id, file_id, "embedding")
            for item in chunks:
                item["metadata"] = build_chunk_metadata(
                    source_file_id=file_id,
                    source_filename=record.get("filename") or record["name"],
                    sheet_name=item.get("metadata", {}).get("sheet_name"),
                    section=item.get("section"),
                    title=item.get("title"),
                    content=item["content"],
                )
            embedding_texts = [embedding_input(item["content"], item["metadata"]) for item in chunks]
            vectors: list[list[float]] = []
            for start in range(0, len(embedding_texts), self.EMBEDDING_BATCH_SIZE):
                vectors.extend(await self.embedding.embed_documents(embedding_texts[start:start + self.EMBEDDING_BATCH_SIZE]))
            for item, vector in zip(chunks, vectors):
                item.update({"embedding": vector, "embedding_provider": self.settings.embedding_provider, "embedding_model": self.settings.embedding_model, "embedding_version": RAG_INDEX_VERSION, "knowledge_base_id": record.get("knowledge_base_id")})
            if not force:
                self.product.set_knowledge_file_status(tenant_id, file_id, "indexing")
            self.product.replace_knowledge_chunks(tenant_id, file_id, chunks)
            self.product.set_knowledge_file_status(tenant_id, file_id, "ready", chunk_count=len(chunks), embedding_provider=self.settings.embedding_provider, embedding_model=self.settings.embedding_model)
        except Exception as exc:
            recovery_status = "ready" if force and previous_status == "ready" else "failed"
            self.product.set_knowledge_file_status(tenant_id, file_id, recovery_status, error_message=str(exc)[:500])
            if raise_errors:
                raise


class KnowledgeRetrievalService:
    def __init__(self, product: ProductStore, settings: Settings) -> None:
        self.product, self.settings = product, settings; self.embedding = embedding_provider_for(settings)
        self.policy = RetrievalConfidencePolicy.from_settings(settings)
        self.answerability_policy = QueryAnswerabilityPolicy(settings.knowledge_query_guard_enabled)
    def search(self, tenant_id: str, query: str, top_k: int = 5) -> list[dict]:
        """Compatibility API for runtime callers that only need accepted chunks."""
        return self.search_outcome(tenant_id, query, top_k).results

    def search_outcome(self, tenant_id: str, query: str, top_k: int = 5) -> RetrievalOutcome:
        guard_reason = self.answerability_policy.pre_retrieval_rejection_reason(query)
        if guard_reason:
            return RetrievalOutcome([], False, "query_guard", guard_reason)
        require_semantic_runtime(self.product, self.settings)
        query_intent = detect_query_intent(query)
        semantic_query = query_embedding_input(query, query_intent)
        if isinstance(self.embedding, OpenAICompatibleEmbeddingProvider):
            query_vector = self.embedding.embed_query_sync(semantic_query)
        elif isinstance(self.embedding, LocalHashEmbeddingProvider):
            query_vector = self.embedding._embed(semantic_query)
        else:
            raise RuntimeError("Embedding Provider 不支持同步检索。")
        try:
            with self.product._store.connection() as conn:
                if self.product._store.is_postgres and self.settings.embedding_provider != "local-hash":
                    literal = "[" + ",".join(f"{value:.10g}" for value in query_vector) + "]"
                    rows = conn.execute(
                        "SELECT c.id,c.file_id,c.content,c.title,c.section,c.page_number,c.metadata,c.embedding_version,"
                        "(1 - (c.embedding <=> ?::vector)) AS vector_score,f.name filename "
                        "FROM knowledge_chunks c JOIN knowledge_files f ON f.id=c.file_id "
                        "WHERE c.tenant_id=? AND f.status='ready' AND c.embedding IS NOT NULL "
                        "AND c.embedding_provider=? AND c.embedding_model=? AND c.embedding_dimension=? "
                        "ORDER BY c.embedding <=> ?::vector LIMIT ?",
                        (literal, tenant_id, self.settings.embedding_provider, self.settings.embedding_model, self.settings.embedding_dimension, literal, max(top_k * 4, 20)),
                    ).fetchall()
                else:
                    rows = conn.execute("SELECT c.id,c.file_id,c.content,c.title,c.section,c.page_number,c.embedding,c.metadata,c.embedding_version,f.name filename FROM knowledge_chunks c JOIN knowledge_files f ON f.id=c.file_id WHERE c.tenant_id=? AND f.status='ready'", (tenant_id,)).fetchall()
        except Exception:
            rows = []
        if not rows:
            # Compatibility for the pre-V1 text records while enterprises migrate
            # their content through the file ingestion pipeline.
            legacy_results = self.product._store.knowledge_search(tenant_id, query, top_k)
            guard_reason = self.answerability_policy.rejection_reason(query, legacy_results)
            if guard_reason:
                return RetrievalOutcome([], False, "query_guard", guard_reason)
            return RetrievalOutcome(legacy_results, bool(legacy_results))
        results=[]
        confidence_reasons: list[str] = []
        for source_row in rows:
            row = dict(source_row)
            if "vector_score" in row:
                semantic = float(row["vector_score"] or 0)
            else:
                vector=json.loads(row["embedding"] or "[]")
                semantic=sum(a*b for a,b in zip(query_vector,vector))
            metadata = metadata_dict(row.get("metadata"))
            keyword = structured_keyword_score(query, title=row["title"] or "", section=row["section"] or "", content=row["content"], metadata=metadata)
            metadata_relevance = calculate_metadata_score(query_intent, metadata)
            score = 0.65 * semantic + 0.35 * keyword + (0.12 * metadata_relevance if metadata_relevance >= 0 else 0.03 * metadata_relevance)
            accepted, rejection_reason = self.policy.decision(vector_score=semantic, keyword_score=keyword, final_score=score)
            if accepted:
                results.append({"chunk_id":row["id"],"file_id":row["file_id"],"filename":row["filename"],"title":row["title"],"content":row["content"][:1200],"section":row["section"],"page":row["page_number"],"query_intent":query_intent.value,"canonical_section":metadata.get("canonical_section"),"metadata_score":round(metadata_relevance,4),"index_version":row.get("embedding_version") or metadata.get("index_version"),"metadata_schema_version":metadata.get("metadata_schema_version"),"vector_score":round(semantic,4),"keyword_score":round(keyword,4),"score":round(score,4),"accepted":True,"rejection_reason":rejection_reason})
            elif rejection_reason:
                confidence_reasons.append(rejection_reason)
        guard_reason = self.answerability_policy.rejection_reason(query, results)
        if guard_reason:
            return RetrievalOutcome([], False, "query_guard", guard_reason)
        results = sorted(results,key=lambda item:item["score"],reverse=True)[:top_k]
        if len(results) >= 2 and results[0]["keyword_score"] <= 0 and results[0]["score"] - results[1]["score"] < self.policy.result_margin:
            return RetrievalOutcome([], False, "within_top_score_margin")
        if not results and confidence_reasons:
            return RetrievalOutcome([], False, confidence_reasons[0])
        return RetrievalOutcome(results, bool(results))
