"""Read-only verification of the real application-exported DOCX, never repair it."""
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from docx import Document

root=Path(sys.argv[1]);path=root/'真实活动方案.docx'
task=json.loads((root/'closurePlanExport.json').read_text(encoding='utf-8'))['final']
expected=task['structured_result']['data'];document=Document(path)
text='\n'.join([p.text for p in document.paragraphs]+[cell.text for table in document.tables for row in table.rows for cell in row.cells])
normalize=lambda value:re.sub(r'\s+','',value)
flat=normalize(text)
for field in ['document_title','activity_theme','activity_time','activity_location','target_audience','invitation_copy']:
    assert normalize(expected[field]) in flat,field
for channel in expected['promotion_channels']:
    assert normalize(channel) in flat,channel
assert len(document.tables)==1
table=document.tables[0]
assert len(table.rows)==len(expected['activity_items'])+1
branded=table.cell(0,3).text=='示意图'
for row in expected['activity_items']:
    # The existing branded template merges phase/name and reserves column 4 for
    # actual illustration images, not image-generation instructions. Do not
    # pretend that ungenerated illustration briefs were exported as pictures.
    fields=['name','description'] if branded else ['phase','name','description','image_requirement']
    for field in fields: assert normalize(row[field]) in flat,(row['name'],field)
if branded:
    assert all(row.cells[3].text=='/' for row in table.rows[1:])
assert document.tables and '/opt/' not in text and 'runtime-data/' not in text
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None
    assert 'word/document.xml' in archive.namelist() and '[Content_Types].xml' in archive.namelist()
    assert not any(name.endswith('vbaProject.bin') for name in archive.namelist())
    for name in archive.namelist():
        if name.endswith(('.xml','.rels')): ET.fromstring(archive.read(name))
result={'task_id':task['id'],'document_opened':True,'parser':'python-docx','ooxml_crc_and_xml_valid':True,'template_content_matches':True,'template':'branded' if branded else 'legacy','illustrations':'NOT_GENERATED; existing branded template displays /' if branded else 'requirements text','table_count':len(document.tables),'activity_items':len(expected['activity_items']),'invitation_complete':True,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'word_desktop_visual_check':'NOT_VERIFIED_NO_OFFICE_RENDERER'}
(root/'closure-docx-content-proof.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
