"""Tenant-aware branded/legacy Activity Plan Word template regression."""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import main
from app.auth import hash_password
from app.document_generator import (
    ActivityPlanContent,
    ActivityPlanDocumentService,
    ActivityPlanDocxRenderer,
    DocumentGenerationError,
    DocumentPresentationOptions,
)
from app.product_store import ProductStore
from app.storage import LocalStorage
from app.store import POCStore


FIXTURES = Path(__file__).parent / "fixtures"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
NS = {"w": W, "wp": WP}
FIXED = lambda: datetime(2026, 9, 24, 8, 30, tzinfo=timezone.utc)


def content(name: str = "d2") -> ActivityPlanContent:
    return ActivityPlanContent.model_validate_json(
        (FIXTURES / f"activity_plan_{name}.json").read_text(encoding="utf-8")
    )


def logo(size=(128, 64), color=(200, 30, 40, 255), *, transparent=True) -> bytes:
    image = Image.new("RGBA", size, color)
    if transparent:
        image.putpixel((0, 0), (color[0], color[1], color[2], 0))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def parts(value: bytes) -> dict[str, bytes]:
    with ZipFile(io.BytesIO(value)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def generated(service: ActivityPlanDocumentService, storage: LocalStorage, tenant="tenant-a",
              presentation: DocumentPresentationOptions | None = None, plan_name="d2") -> bytes:
    result = service.create(tenant, "user-a", content(plan_name), presentation or DocumentPresentationOptions())
    return storage.get(result.storage_key)


@pytest.fixture
def setup(tmp_path: Path):
    storage = LocalStorage(tmp_path / "objects")
    configs = {
        "tenant-a": {"brand_name": "甲企业", "brand_logo": "brand-assets/tenant-a/logos/a.png"},
        "tenant-b": {"brand_name": "乙企业", "brand_logo": "brand-assets/tenant-b/logos/b.png"},
        "tenant-c": {"brand_name": "丙企业"},
    }
    storage.put(configs["tenant-a"]["brand_logo"], logo(), "image/png")
    storage.put(configs["tenant-b"]["brand_logo"], logo(color=(30, 40, 200, 255)), "image/png")
    service = ActivityPlanDocumentService(storage, clock=FIXED, brand_provider=lambda tenant: configs[tenant])
    return service, storage, configs


def test_w1_w2_w3_w4_w7_w8_w10_w11_branded_structure_and_trace(setup, caplog):
    service, storage, configs = setup
    with caplog.at_level(logging.INFO, logger="app.document_generator"):
        value = generated(service, storage, presentation=DocumentPresentationOptions(brand_name="伪造企业"),
                          plan_name="d1")
    assert service.resolve_activity_plan_template("tenant-a").variant == "branded"
    assert '"template_variant": "branded"' in caplog.text
    assert '"fallback_reason": null' in caplog.text
    assert configs["tenant-a"]["brand_logo"] not in caplog.text
    doc = Document(io.BytesIO(value))
    section = doc.sections[0]
    assert abs(section.page_width.cm - 21) < .01
    assert abs(section.page_height.cm - 29.7) < .01
    assert abs(section.top_margin.cm - 2.54) < .01
    assert abs(section.bottom_margin.cm - 2.54) < .01
    assert abs(section.left_margin.cm - 3.175) < .01
    assert abs(section.right_margin.cm - 3.175) < .01
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.styles["Normal"].font.name == "FangSong"
    assert doc.styles["Normal"].font.size.pt == 12
    assert doc.styles["Title"].font.size.pt == 20
    assert any(p.text.startswith("附件一：") and p.paragraph_format.page_break_before for p in doc.paragraphs)

    payload = parts(value)
    root = ElementTree.fromstring(payload["word/document.xml"])
    header = ElementTree.fromstring(payload["word/header1.xml"])
    assert "甲企业-活动方案  请勿外传" in payload["word/header1.xml"].decode()
    assert "伪造企业" not in b"".join(payload.values()).decode(errors="ignore")
    anchors = header.findall(".//wp:anchor", NS)
    assert len(anchors) == 2
    assert [anchor.get("behindDoc") for anchor in anchors] == ["0", "1"]
    for anchor, (expected_x, expected_y) in zip(anchors, ((28800, 64800), (1143000, 2710800))):
        x = int(anchor.find("wp:positionH/wp:posOffset", NS).text)
        y = int(anchor.find("wp:positionV/wp:posOffset", NS).text)
        assert abs(x - expected_x) <= 100
        assert abs(y - expected_y) <= 100
    alpha_peaks = []
    for name, image_bytes in payload.items():
        if name.startswith("word/media/") and name.endswith(".png"):
            with Image.open(io.BytesIO(image_bytes)) as image:
                alpha_peaks.append(image.convert("RGBA").getchannel("A").getextrema()[1])
    assert 153 in alpha_peaks and 26 in alpha_peaks
    assert b" PAGE " in payload["word/footer1.xml"]
    assert "第".encode() not in payload["word/footer1.xml"]
    table = root.find(".//w:tbl", NS)
    assert table is not None
    widths = [int(node.get(f"{{{W}}}w")) for node in table.findall("w:tblGrid/w:gridCol", NS)]
    assert all(abs(actual - target) <= 2 for actual, target in zip(widths, [896, 1298, 4462, 2279]))
    rows = table.findall("w:tr", NS)
    assert len(rows) == 5
    assert rows[0].find("w:trPr/w:tblHeader", NS) is not None
    for cell in rows[0].findall("w:tc", NS):
        assert cell.find("w:tcPr/w:shd", NS).get(f"{{{W}}}fill") == "D8D8D8"
    for row in rows[1:]:
        cells = row.findall("w:tc", NS)
        assert len(cells) == 3
        assert cells[0].find("w:tcPr/w:gridSpan", NS).get(f"{{{W}}}val") == "2"
        assert all(cell.find("w:tcPr/w:shd", NS) is None for cell in cells)
        assert row.find("w:trPr/w:trHeight", NS) is None
    assert any("".join(t.text or "" for t in cell.findall(".//w:t", NS)) == "/"
               for row in rows[1:] for cell in row.findall("w:tc", NS))
    assert all(border.get(f"{{{W}}}color") == "000000" and border.get(f"{{{W}}}sz") == "4"
               for border in table.find("w:tblPr/w:tblBorders", NS))


@pytest.mark.parametrize("size", [(400, 60), (60, 400)])
def test_w5_w6_logo_aspect_ratio_is_preserved(setup, size):
    service, storage, configs = setup
    storage.put(configs["tenant-a"]["brand_logo"], logo(size=size), "image/png")
    payload = parts(generated(service, storage))
    header = ElementTree.fromstring(payload["word/header1.xml"])
    for anchor in header.findall(".//wp:anchor", NS):
        extent = anchor.find("wp:extent", NS)
        rendered_ratio = int(extent.get("cx")) / int(extent.get("cy"))
        assert abs(rendered_ratio - size[0] / size[1]) < .002


def test_w9_real_image_keeps_ratio(setup):
    service, storage, _ = setup
    image = Image.new("RGB", (180, 90), "green")
    output = io.BytesIO()
    image.save(output, format="JPEG")
    key = "brand-assets/tenant-a/activity-plan-images/scene.jpg"
    storage.put(key, output.getvalue(), "image/jpeg")
    doc = Document(io.BytesIO(generated(
        service, storage, presentation=DocumentPresentationOptions(image_paths={0: key})
    )))
    assert len(doc.inline_shapes) == 1
    assert abs(doc.inline_shapes[0].width / doc.inline_shapes[0].height - 2) < .002


def test_branded_image_paths_are_tenant_scoped(setup):
    service, storage, _ = setup
    with pytest.raises(DocumentGenerationError, match="示意图"):
        generated(service, storage, presentation=DocumentPresentationOptions(
            image_paths={0: "brand-assets/tenant-b/activity-plan-images/scene.jpg"}
        ))


def test_w12_next_generation_uses_replaced_logo(setup):
    service, storage, configs = setup
    first = parts(generated(service, storage))
    new_key = "brand-assets/tenant-a/logos/replacement.png"
    storage.put(new_key, logo(color=(10, 180, 40, 255)), "image/png")
    configs["tenant-a"]["brand_logo"] = new_key
    second = parts(generated(service, storage))
    first_media = {value for name, value in first.items() if name.startswith("word/media/")}
    second_media = {value for name, value in second.items() if name.startswith("word/media/")}
    assert first_media.isdisjoint(second_media)


def test_w13_f6_tenant_isolation_and_branded_legacy_coexistence(setup):
    service, storage, _ = setup
    a = parts(generated(service, storage, "tenant-a"))
    b = parts(generated(service, storage, "tenant-b"))
    c = parts(generated(service, storage, "tenant-c"))
    assert service.resolve_activity_plan_template("tenant-a").variant == "branded"
    assert service.resolve_activity_plan_template("tenant-b").variant == "branded"
    assert service.resolve_activity_plan_template("tenant-c").variant == "legacy"
    assert "甲企业" in a["word/header1.xml"].decode()
    assert "乙企业" in b["word/header1.xml"].decode()
    assert "甲企业" not in b["word/header1.xml"].decode()
    assert a["word/header1.xml"] != b["word/header1.xml"]
    assert b"wp:anchor" not in c["word/header1.xml"]


def test_f1_f5_missing_logo_uses_unmodified_legacy_renderer(setup, caplog):
    service, storage, _ = setup
    with caplog.at_level(logging.INFO, logger="app.document_generator"):
        value = generated(service, storage, "tenant-c")
    selection = service.resolve_activity_plan_template("tenant-c")
    assert (selection.variant, selection.fallback_reason) == ("legacy", "logo_missing")
    assert '"template_variant": "legacy"' in caplog.text
    assert '"fallback_reason": "logo_missing"' in caplog.text
    expected = ActivityPlanDocxRenderer(FIXED).render(content(), DocumentPresentationOptions())
    actual_parts, expected_parts = parts(value), parts(expected)
    for name in ("word/document.xml", "word/styles.xml", "word/header1.xml", "word/footer1.xml"):
        assert actual_parts[name] == expected_parts[name]
    assert Document(io.BytesIO(value)).styles["Title"].font.size.pt == 26


@pytest.mark.parametrize("invalid", [logo(transparent=False), b"not a png"])
def test_f2_invalid_logo_falls_back(setup, invalid):
    service, storage, configs = setup
    storage.put(configs["tenant-a"]["brand_logo"], invalid, "image/png")
    selection = service.resolve_activity_plan_template("tenant-a")
    assert (selection.variant, selection.fallback_reason) == ("legacy", "logo_invalid")
    assert Document(io.BytesIO(generated(service, storage))).styles["Title"].font.size.pt == 26


def test_f3_unreadable_logo_falls_back(setup, monkeypatch):
    service, storage, configs = setup
    original_get = storage.get
    def fail_logo(key):
        if key == configs["tenant-a"]["brand_logo"]:
            raise OSError("storage unavailable")
        return original_get(key)
    monkeypatch.setattr(storage, "get", fail_logo)
    selection = service.resolve_activity_plan_template("tenant-a")
    assert (selection.variant, selection.fallback_reason) == ("legacy", "logo_unreadable")
    assert generated(service, storage).startswith(b"PK")


def test_f4_ownership_mismatch_falls_back(setup):
    service, storage, configs = setup
    configs["tenant-a"]["brand_logo"] = configs["tenant-b"]["brand_logo"]
    selection = service.resolve_activity_plan_template("tenant-a")
    assert (selection.variant, selection.fallback_reason) == ("legacy", "ownership_invalid")
    assert generated(service, storage).startswith(b"PK")


def test_config_read_failure_preserves_legacy_export(setup):
    _, storage, _ = setup
    def unavailable(_tenant):
        raise OSError("config unavailable")
    service = ActivityPlanDocumentService(storage, clock=FIXED, brand_provider=unavailable)
    assert service.resolve_activity_plan_template("tenant-a").fallback_reason == "logo_unreadable"
    assert generated(service, storage).startswith(b"PK")


def test_existing_post_get_api_serves_branded_docx(tmp_path: Path, monkeypatch):
    store = POCStore(tmp_path / "api.db")
    store.seed_demo_data()
    products = ProductStore(store)
    products.initialize()
    products.create_user("tenant-a", "admin@tenant-a.test", hash_password("ChangeMe!2026"),
                         "Tenant A Admin", "enterprise_admin")
    storage = LocalStorage(tmp_path / "api-objects")
    key = "brand-assets/tenant-a/logos/api.png"
    storage.put(key, logo(), "image/png")
    products.set_brand_logo("tenant-a", key, {
        "asset_id": "api", "filename": "api.png", "content_type": "image/png",
        "width": 128, "height": 64, "has_transparency": True,
    })
    monkeypatch.setattr(main, "store", store)
    monkeypatch.setattr(main, "product_store", products)
    monkeypatch.setattr(main, "document_service", ActivityPlanDocumentService(
        storage, brand_provider=lambda tenant: store.enterprise_config(tenant),
    ))
    client = TestClient(main.app)
    try:
        assert client.post("/api/v1/auth/login", json={
            "account": "admin@tenant-a.test", "password": "ChangeMe!2026",
        }).status_code == 200
        response = client.post("/api/v1/documents/activity-plan", json={
            "content": content("d1").model_dump(), "presentation": {},
        })
        assert response.status_code == 201
        result = response.json()
        assert set(result) == {"document_id", "filename", "content_type", "download_url"}
        download = client.get(result["download_url"])
        assert download.status_code == 200
        assert "启明教育-活动方案" in parts(download.content)["word/header1.xml"].decode()
    finally:
        client.close()
