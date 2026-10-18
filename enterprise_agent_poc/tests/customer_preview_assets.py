"""Generate clearly marked, non-customer documents for real Preview ingestion."""
import sys
from pathlib import Path
from docx import Document
from PIL import Image, ImageDraw

root = Path(sys.argv[1]); root.mkdir(parents=True,exist_ok=True)
document = Document()
document.add_heading('星河体验教育测试课程',0)
document.add_paragraph('此文档仅为独立Preview验收材料，不属于真实客户。')
document.add_heading('天文观察课',1)
document.add_paragraph('课程面向6至9岁儿童与家长，时间为周三19点，地点为测试体验教室。课程包含星图手工和基础观察体验。价格尚未确认，不承诺学习效果。')
document.add_heading('企业规则',1)
document.add_paragraph('品牌名称为星河体验教育。宣传语保持清晰可信，不虚构价格、师资及名额。')
document.save(root/'星河体验课程_真实Word上传.docx')
image=Image.new('RGBA',(128,128),(0,0,0,0)); draw=ImageDraw.Draw(image);draw.ellipse((8,8,120,120),fill=(27,78,145,255));draw.text((57,54),'P',fill='white')
image.save(root/'preview-logo.png')
