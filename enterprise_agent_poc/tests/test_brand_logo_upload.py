from __future__ import annotations

import io
import tomllib
from pathlib import Path
from zipfile import ZipFile

import PIL
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import main
from app.auth import hash_password
from app.brand_logo import (
    BrandLogoError, BrandLogoService, MAX_LOGO_BYTES, PNG_SIGNATURE, inspect_transparent_png,
    validate_brand_logo_key,
)
from app.document_generator import (
    ActivityPlanContent, ActivityPlanDocumentService, DocumentPresentationOptions,
)
from app.product_store import ProductStore
from app.storage import LocalStorage
from app.store import POCStore


ROOT = Path(__file__).resolve().parents[1]
ENDPOINT = "/api/v1/enterprise-config/brand-logo"


def png(*, mode: str = "RGBA", size: tuple[int, int] = (64, 32), translucent: bool = True) -> bytes:
    image = Image.new(mode, size, (240, 40, 40, 255) if mode == "RGBA" else (240, 40, 40))
    if mode == "RGBA" and translucent:
        image.putpixel((0, 0), (240, 40, 40, 0))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


@pytest.fixture
def isolated_api(tmp_path: Path, monkeypatch):
    store = POCStore(tmp_path / "brand-logo.db")
    store.seed_demo_data()
    products = ProductStore(store)
    products.initialize()
    for tenant in ("tenant-a", "tenant-b"):
        products.create_user(tenant, f"admin@{tenant}.test", hash_password("ChangeMe!2026"),
                             f"{tenant} Admin", "enterprise_admin")
        products.create_user(tenant, f"member@{tenant}.test", hash_password("ChangeMe!2026"),
                             f"{tenant} Member", "member")
    storage = LocalStorage(tmp_path / "objects")
    monkeypatch.setattr(main, "store", store)
    monkeypatch.setattr(main, "product_store", products)
    monkeypatch.setattr(main, "brand_logo_service", BrandLogoService(storage, products))
    monkeypatch.setattr(main, "document_service", ActivityPlanDocumentService(
        storage,
    ))
    return TestClient(main.app), store, products, storage


def login(client: TestClient, tenant: str = "tenant-a", role: str = "admin") -> None:
    client.post("/api/v1/auth/logout")
    response = client.post("/api/v1/auth/login", json={
        "account": f"{role}@{tenant}.test", "password": "ChangeMe!2026",
    })
    assert response.status_code == 200


def upload(client: TestClient, data: bytes, filename: str = "logo.png"):
    return client.post(ENDPOINT, files={"file": (filename, data, "image/png")})


def test_l1_valid_transparent_png_upload_preview_and_config(isolated_api):
    client, store, _, storage = isolated_api
    login(client)
    content = png()
    response = upload(client, content)
    assert response.status_code == 201
    metadata = response.json()["brand_logo"]
    assert metadata["filename"] == "logo.png"
    assert metadata["content_type"] == "image/png"
    assert (metadata["width"], metadata["height"]) == (64, 32)
    assert metadata["has_transparency"] is True
    assert metadata["download_url"] == ENDPOINT
    key = store.enterprise_config("tenant-a")["brand_logo"]
    assert key.startswith("brand-assets/tenant-a/logos/") and key.endswith(".png")
    assert storage.get(key) == content
    public = client.get("/api/v1/enterprise-config")
    assert public.status_code == 200 and public.json()["brand_logo"] == metadata
    assert key not in public.text
    preview = client.get(metadata["download_url"])
    assert preview.status_code == 200 and preview.content == content
    assert preview.headers["cache-control"] == "private, no-store"


@pytest.mark.parametrize("data", [png(mode="RGB"), png(translucent=False)])
def test_l2_opaque_png_rejected_even_when_rgba(data: bytes):
    with pytest.raises(BrandLogoError, match="LOGO_TRANSPARENCY_REQUIRED"):
        inspect_transparent_png(data, filename="logo.png")


def test_l3_fake_png_and_unsupported_extension_are_blocked(isolated_api):
    client, _, _, _ = isolated_api
    login(client)
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(buffer, format="JPEG")
    assert upload(client, buffer.getvalue()).json()["detail"] == "LOGO_FORMAT_NOT_SUPPORTED"
    assert upload(client, png(), "logo.jpg").json()["detail"] == "LOGO_FORMAT_NOT_SUPPORTED"


def test_l4_corrupted_png_is_blocked(isolated_api):
    client, _, _, _ = isolated_api
    login(client)
    damaged = bytearray(png())
    damaged[-8] ^= 0xFF  # Break the final chunk CRC while keeping a valid PNG signature.
    response = upload(client, bytes(damaged))
    assert response.status_code == 422 and response.json()["detail"] == "LOGO_INVALID_PNG"


def test_l5_oversized_file_is_blocked(isolated_api):
    client, _, _, _ = isolated_api
    login(client)
    response = upload(client, PNG_SIGNATURE + b"0" * MAX_LOGO_BYTES)
    assert response.status_code == 413 and response.json()["detail"] == "LOGO_TOO_LARGE"


def test_l6_oversized_dimensions_are_blocked(isolated_api):
    client, _, _, _ = isolated_api
    login(client)
    response = upload(client, png(size=(4097, 1)))
    assert response.status_code == 422 and response.json()["detail"] == "LOGO_DIMENSIONS_TOO_LARGE"


def test_l7_tenant_isolation_and_no_arbitrary_config_reference(isolated_api):
    client, store, _, _ = isolated_api
    login(client, "tenant-a")
    assert upload(client, png()).status_code == 201
    tenant_a_key = store.enterprise_config("tenant-a")["brand_logo"]
    rejected = client.put("/api/v1/enterprise-config", json={"payload": {"brand_logo": tenant_a_key}})
    assert rejected.status_code == 422 and rejected.json()["detail"] == "BRAND_LOGO_UPLOAD_REQUIRED"
    rejected = client.put("/api/v1/enterprise-config", json={"payload": {"brand_logo_metadata": {"storage_key": tenant_a_key}}})
    assert rejected.status_code == 422
    login(client, "tenant-b")
    assert client.get(ENDPOINT).status_code == 404
    assert client.get("/api/v1/enterprise-config").json()["brand_logo"] is None
    assert "brand_logo" not in store.enterprise_config("tenant-b")
    with pytest.raises(BrandLogoError, match="LOGO_PERMISSION_DENIED"):
        validate_brand_logo_key(tenant_a_key, "tenant-b")
    with store.connection() as conn:
        import json
        config = store.enterprise_config("tenant-b")
        config["brand_logo"] = tenant_a_key
        conn.execute("UPDATE enterprise_configs SET payload=? WHERE tenant_id=?",
                     (json.dumps(config), "tenant-b"))
    assert client.get(ENDPOINT).json()["detail"] == "LOGO_PERMISSION_DENIED"
    login(client, "tenant-a", "member")
    assert upload(client, png()).json()["detail"] == "LOGO_PERMISSION_DENIED"
    assert client.get(ENDPOINT).json()["detail"] == "LOGO_PERMISSION_DENIED"


def test_l8_failed_validation_and_storage_preserve_old_logo(isolated_api, monkeypatch):
    client, store, products, storage = isolated_api
    login(client)
    assert upload(client, png()).status_code == 201
    before = store.enterprise_config("tenant-a")["brand_logo"]
    assert upload(client, png(mode="RGB")).status_code == 422
    assert store.enterprise_config("tenant-a")["brand_logo"] == before
    original_put = storage.put
    def fail_put(*_args):
        raise OSError("storage down")
    monkeypatch.setattr(storage, "put", fail_put)
    failed = upload(client, png())
    assert failed.status_code == 503 and failed.json()["detail"] == "LOGO_STORAGE_FAILED"
    assert store.enterprise_config("tenant-a")["brand_logo"] == before
    monkeypatch.setattr(storage, "put", original_put)
    def fail_update(*_args):
        raise RuntimeError("config transaction failed")
    monkeypatch.setattr(products, "set_brand_logo", fail_update)
    failed = upload(client, png())
    assert failed.status_code == 503 and failed.json()["detail"] == "LOGO_CONFIG_UPDATE_FAILED"
    assert store.enterprise_config("tenant-a")["brand_logo"] == before


def test_l9_valid_replacement_switches_pointer_and_keeps_old_asset(isolated_api):
    client, store, _, storage = isolated_api
    login(client)
    assert upload(client, png()).status_code == 201
    first = store.enterprise_config("tenant-a")["brand_logo"]
    replacement = png(size=(128, 64))
    assert upload(client, replacement).status_code == 201
    second = store.enterprise_config("tenant-a")["brand_logo"]
    assert first != second and storage.get(first) == png()
    assert storage.get(second) == replacement
    assert client.get(ENDPOINT).content == replacement


def test_a1_a2_missing_logo_keeps_existing_word_export(isolated_api):
    client, store, _, storage = isolated_api
    login(client)
    assert client.get("/api/v1/enterprise-config").json()["brand_logo"] is None
    assert "brand_logo" not in store.enterprise_config("tenant-a")
    content = ActivityPlanContent.model_validate_json(
        (ROOT / "tests" / "fixtures" / "activity_plan_d2.json").read_text(encoding="utf-8")
    )
    generated = main.document_service.create("tenant-a", "admin-a", content, DocumentPresentationOptions())
    value = storage.get(generated.storage_key)
    assert value.startswith(b"PK")
    with ZipFile(io.BytesIO(value)) as archive:
        document = Document(io.BytesIO(value))
        assert document.paragraphs
        assert "word/document.xml" in archive.namelist()
    assert b"BRAND_LOGO_REQUIRED_FOR_ACTIVITY_PLAN_TEMPLATE" not in value


def test_a12_existing_word_export_with_logo_remains_optional(isolated_api):
    client, store, _, storage = isolated_api
    login(client)
    assert upload(client, png()).status_code == 201
    assert store.enterprise_config("tenant-a")["brand_logo"]
    content = ActivityPlanContent.model_validate_json(
        (ROOT / "tests" / "fixtures" / "activity_plan_d2.json").read_text(encoding="utf-8")
    )
    generated = main.document_service.create("tenant-a", "admin-a", content, DocumentPresentationOptions())
    value = storage.get(generated.storage_key)
    with ZipFile(io.BytesIO(value)) as archive:
        assert "word/document.xml" in archive.namelist()
    assert Document(io.BytesIO(value)).paragraphs


def test_a11_formal_dependency_and_release_preflight():
    assert PIL.__version__
    project = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    assert "Pillow>=10,<13" in project["project"]["dependencies"]
    switch = (ROOT / "deploy" / "release_switch.sh").read_text(encoding="utf-8")
    assert "from PIL import Image; import PIL" in switch
    installer = (ROOT / "deploy" / "prepare_runtime_dependencies.sh").read_text(encoding="utf-8")
    assert '"Pillow>=10,<13"' in installer
    assert '"$runtime_venv/bin/pip" install' in installer
    assert '"$runtime_venv/bin/pip" check' in installer


def test_upload_requires_a_file(isolated_api):
    client, _, _, _ = isolated_api
    login(client)
    response = client.post(ENDPOINT)
    assert response.status_code == 400 and response.json()["detail"] == "LOGO_FILE_REQUIRED"
